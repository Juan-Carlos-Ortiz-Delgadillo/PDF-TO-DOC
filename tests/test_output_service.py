from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.core.exceptions import OutputValidationError
from app.services.output_service import OutputService


def _write_docx(path: Path, text: str) -> None:
    document = Document()
    document.add_paragraph(text)
    document.save(path)


def test_output_service_creates_and_cleans_temporary_files(tmp_path: Path) -> None:
    service = OutputService()
    destination = tmp_path / "result.docx"

    with service.temporary_output(destination) as temporary_path:
        assert temporary_path.exists()
        _write_docx(temporary_path, "temporary")
    assert not temporary_path.exists()

    created = service.create_temporary_output(
        tmp_path / "nested" / "result.docx",
        create_directory=True,
    )
    assert created.parent.is_dir()
    service.cleanup_temporary(created)
    assert not created.exists()


def test_output_service_validates_and_publishes_docx(tmp_path: Path) -> None:
    service = OutputService()
    destination = tmp_path / "result.docx"
    temporary_path = service.create_temporary_output(destination)
    _write_docx(temporary_path, "published")

    assert service.validate_docx(temporary_path) == temporary_path
    assert service.publish(temporary_path, destination) == destination
    assert destination.exists()
    assert not temporary_path.exists()


def test_output_service_requires_confirmation_to_replace(tmp_path: Path) -> None:
    service = OutputService()
    destination = tmp_path / "result.docx"
    _write_docx(destination, "old")
    temporary_path = service.create_temporary_output(
        destination,
        overwrite_confirmed=True,
    )
    _write_docx(temporary_path, "new")

    assert (
        service.publish(
            temporary_path,
            destination,
            overwrite_confirmed=True,
        )
        == destination
    )
    assert "new" in "\n".join(paragraph.text for paragraph in Document(destination).paragraphs)


def test_output_service_rejects_wrong_destinations_and_invalid_documents(
    tmp_path: Path,
) -> None:
    service = OutputService()
    with pytest.raises(ValueError):
        OutputService(minimum_size_bytes=3)
    with pytest.raises(OutputValidationError):
        service.ensure_output_directory(tmp_path / "missing")
    with pytest.raises(OutputValidationError):
        service.validate_destination(tmp_path / "bad.pdf")

    invalid_docx = tmp_path / "invalid.docx"
    invalid_docx.write_text("not a docx", encoding="utf-8")
    with pytest.raises(OutputValidationError):
        service.validate_docx(invalid_docx)

    with pytest.raises(OutputValidationError):
        service.validate_docx(tmp_path / "missing.docx")
